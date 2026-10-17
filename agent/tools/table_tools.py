"""Bounded, read-only CSV/XLSX tools without model downloads or domain rules."""
import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path
import json
import zipfile
from agent.tools.project_tools import ProjectTools, SKIP_DIRS, SKIP_FILES


class TableTools(ProjectTools):
    MAX_ROWS = 5000
    MAX_COLUMNS = 100

    def _table(self, relative_path):
        requested = Path(relative_path)
        if requested.is_absolute() or ".." in requested.parts:
            raise ValueError("Use a relative table path inside the project root.")
        path = (self.root / requested).resolve()
        if not path.is_relative_to(self.root):
            raise ValueError("Table must be inside the project root.")
        parts = path.relative_to(self.root).parts
        if any(part in SKIP_DIRS or part.startswith(".env") or part in SKIP_FILES for part in parts):
            raise ValueError("Excluded directory or credential path.")
        if path.suffix.lower() not in {".csv", ".xlsx"} or not path.is_file():
            raise ValueError("Choose an existing CSV or XLSX file.")
        if path.stat().st_size > 5_000_000:
            raise ValueError("Table exceeds the 5 MB file limit.")
        if path.suffix.lower() == ".xlsx":
            with zipfile.ZipFile(path) as archive:
                if sum(item.file_size for item in archive.infolist()) > 50_000_000:
                    raise ValueError("Workbook exceeds the expanded size limit.")
        return path

    def list_tables(self, limit=50):
        if type(limit) is not int or not 1 <= limit <= 50:
            raise ValueError("limit must be between 1 and 50.")
        def walk(folder):
            for path in sorted(folder.iterdir()):
                if path.is_symlink() or path.name.startswith(".env") or path.name in SKIP_FILES:
                    continue
                if path.is_dir():
                    if path.name not in SKIP_DIRS:
                        yield from walk(path)
                elif path.suffix.lower() in {".csv", ".xlsx"}:
                    yield path.relative_to(self.root).as_posix()
        files=[]
        for name in walk(self.root):
            if len(files) == limit:
                return {"files":files,"truncated":True}
            files.append(name)
        return {"files":files,"truncated":False}

    def search_table(self, relative_path, query=None, sheet=None, header_row=1,
                     column=None, operator="contains", value=None, limit=10):
        if type(header_row) is not int or not 1 <= header_row <= 50:
            raise ValueError("header_row must be between 1 and 50.")
        if type(limit) is not int or not 1 <= limit <= 20:
            raise ValueError("limit must be between 1 and 20.")
        if operator not in {"contains","eq","gt","ge","lt","le"}:
            raise ValueError("Unsupported filter operator.")
        if (column is None) != (value is None) or (column is None and operator != "contains"):
            raise ValueError("Supply column and value together for a filter.")
        path=self._table(relative_path)
        workbook=None
        handle=None
        try:
            if path.suffix.lower()==".xlsx":
                from openpyxl import load_workbook
                workbook=load_workbook(path,read_only=True,data_only=True,keep_links=False)
                if sheet is None and len(workbook.sheetnames)>1:
                    return {"file":relative_path,"sheets":workbook.sheetnames,
                            "error":"Choose a sheet and call search_table again."}
                selected=sheet if sheet is not None else workbook.sheetnames[0]
                if selected not in workbook.sheetnames:
                    return {"error":"Sheet not found.","sheets":workbook.sheetnames}
                ws=workbook[selected]
                if ws.max_column and ws.max_column>self.MAX_COLUMNS:
                    raise ValueError("Sheet exceeds the 100-column limit.")
                rows=ws.iter_rows(min_row=header_row,max_row=header_row+self.MAX_ROWS+1,
                                  max_col=self.MAX_COLUMNS,values_only=True)
            else:
                if sheet is not None:
                    raise ValueError("CSV files do not have sheets.")
                selected=None
                handle=path.open(encoding="utf-8-sig",newline="")
                rows=iter(csv.reader(handle))
                for _ in range(header_row-1):
                    next(rows,None)
            headers=[str(v).strip() if v is not None else "" for v in next(rows,())]
            while headers and not headers[-1]:
                headers.pop()
            if not headers or len(headers)>self.MAX_COLUMNS or not all(headers) or len(set(headers))!=len(headers):
                return {"error":"Choose a header_row with nonempty unique column names (maximum 100)."}
            if column is not None and column not in headers:
                return {"error":"Filter column not found.","columns":headers}
            expected=None
            if column is not None and operator in {"gt","ge","lt","le"}:
                try:
                    expected=Decimal(value)
                    if not expected.is_finite(): raise InvalidOperation
                except InvalidOperation:
                    raise ValueError("Numeric comparison requires a finite numeric value.")
            terms=(query or "").casefold().split()
            results=[]; scanned=0; truncated=False; output_size=0
            for row_number,values in enumerate(rows,header_row+1):
                if scanned>=self.MAX_ROWS:
                    truncated=True; break
                scanned+=1
                if len(values)>self.MAX_COLUMNS:
                    raise ValueError("Row exceeds the 100-column limit.")
                record={h:(str(values[i]) if i<len(values) and values[i] is not None else "") for i,h in enumerate(headers)}
                if not any(record.values()): continue
                haystack=" ".join(record.values()).casefold()
                if not all(term in haystack for term in terms): continue
                if column is not None:
                    actual=record[column]
                    if operator=="contains" and value.casefold() not in actual.casefold(): continue
                    if operator=="eq" and actual.casefold()!=value.casefold(): continue
                    if expected is not None:
                        try:
                            number=Decimal(actual)
                            if not number.is_finite(): continue
                        except InvalidOperation:
                            continue
                        matched={"gt":number>expected,"ge":number>=expected,"lt":number<expected,"le":number<=expected}[operator]
                        if not matched: continue
                if len(results)>=limit:
                    truncated=True; break
                item={"row":row_number,"values":{k:v[:500] for k,v in record.items()},
                      "cells_truncated":any(len(v)>500 for v in record.values())}
                size=len(json.dumps(item,ensure_ascii=False))
                if output_size+size>5500:
                    truncated=True; break
                results.append(item); output_size+=size
            return {"file":path.relative_to(self.root).as_posix(),"sheet":selected,
                    "header_row":header_row,"columns":headers,"rows":results,
                    "rows_scanned":scanned,"truncated":truncated,
                    "note":"Bounded row search, not a full-table count. XLSX uses cached values; formulas are not executed."}
        finally:
            if workbook is not None: workbook.close()
            if handle is not None: handle.close()
