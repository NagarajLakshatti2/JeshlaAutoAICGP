"""Explicit document ingestion; importing this module never builds indexes."""
from pathlib import Path
if __package__:
    from .paths import RAG_DIR, DATA_FOLDER
else:
    from paths import RAG_DIR, DATA_FOLDER


def create_vector_db(data_folder=None, index_folder=None):
    """Build trusted local indexes on explicit request; defaults match retrieval."""
    import pickle

    from pypdf import PdfReader
    from docx import Document as DocxDocument
    import csv

    from langchain_core.documents import Document
    from langchain_huggingface import HuggingFaceEmbeddings
    from langchain_community.vectorstores import FAISS
    from langchain_text_splitters import RecursiveCharacterTextSplitter

    from rank_bm25 import BM25Okapi

    from pathlib import Path

    data_folder = Path(data_folder).expanduser().resolve() if data_folder is not None else DATA_FOLDER


    if not data_folder.is_dir():
        raise ValueError("Choose an existing input data directory; no indexes were changed.")

    pdf_files = list(data_folder.glob("*.pdf"))
    txt_files = list(data_folder.glob("*.txt"))
    docx_files = list(data_folder.glob("*.docx"))
    csv_files = list(data_folder.glob("*.csv"))


    print(f"PDF files found: {len(pdf_files)}")
    print(f"TXT files found: {len(txt_files)}")
    print(f"DOCX files found: {len(docx_files)}")
    print(f"CSV files found: {len(csv_files)}")


    # --------------------------------------------------
    # Configuration
    # --------------------------------------------------

    input_folder = data_folder

    output_folder = Path(index_folder).expanduser().resolve() if index_folder is not None else RAG_DIR
    output_folder.mkdir(parents=True, exist_ok=True)
    FAISS_DB_FOLDER = output_folder / "faiss_db"
    BM25_FILE = output_folder / "bm25.pkl"
    CHUNKS_FILE = output_folder / "chunks.pkl"


    # --------------------------------------------------
    # Find PDF files
    # --------------------------------------------------

    pdf_files = list(input_folder.glob("*.pdf"))

    print(f"PDF files found: {len(pdf_files)}")


    # --------------------------------------------------
    # Read PDF files
    # --------------------------------------------------

    documents = []

    # ==============================
    # PDF FILES
    # ==============================

    for pdf_path in pdf_files:
        print(f"\nReading PDF: {pdf_path.name}")

        reader = PdfReader(pdf_path)

        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""

            if text.strip():
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": pdf_path.name,
                            "page": page_number,
                            "file_type": "pdf"
                        }
                    )
                )


    # ==============================
    # TXT FILES
    # ==============================

    txt_files = list(data_folder.glob("*.txt"))

    for txt_path in txt_files:
        print(f"\nReading TXT: {txt_path.name}")

        text = txt_path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": txt_path.name,
                        "page": 1,
                        "file_type": "txt"
                    }
                )
            )

    # DOCX
    for docx_path in docx_files:
        print(f"\nReading DOCX: {docx_path.name}")

        docx_document = DocxDocument(docx_path)

        text = "\n".join(
            paragraph.text
            for paragraph in docx_document.paragraphs
            if paragraph.text.strip()
        )

        if text.strip():
            documents.append(
                Document(
                    page_content=text,
                    metadata={
                        "source": docx_path.name,
                        "page": 1,
                        "file_type": "docx"
                    }
                )
            )

    # ============================================================
    # LOAD CSV FILES
    # ============================================================

    for csv_path in csv_files:

        print(f"\nReading CSV: {csv_path.name}")

        with open(
            csv_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.DictReader(file)

            for row_number, row in enumerate(reader, start=2):

                row_text = "\n".join(
                    f"{column}: {value}"
                    for column, value in row.items()
                    if value is not None and str(value).strip()
                )

                if row_text.strip():

                    documents.append(
                        Document(
                            page_content=row_text,
                            metadata={
                                "source": csv_path.name,
                                "page": row_number,
                                "file_type": "csv",
                                "row": row_number
                            }
                        )
                    )

    if not documents:
        raise ValueError("No readable PDF, TXT, DOCX or CSV content found; existing indexes were not changed.")

    print(f"Total documents loaded: {len(documents)}")

    # --------------------------------------------------
    # Chunking
    # --------------------------------------------------

    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=500,
        chunk_overlap=50
    )

    chunks = text_splitter.split_documents(documents)

    print(f"Total chunks: {len(chunks)}")


    # --------------------------------------------------
    # Create BM25 index
    # --------------------------------------------------

    print("\nCreating BM25 index...")

    tokenized_chunks = [
        chunk.page_content.lower().split()
        for chunk in chunks
    ]

    bm25 = BM25Okapi(tokenized_chunks)


    # --------------------------------------------------
    # Save BM25 index
    # --------------------------------------------------

    with open(BM25_FILE, "wb") as file:

        pickle.dump(
            bm25,
            file
        )


    print(
        f"BM25 database saved: {BM25_FILE}"
    )


    # --------------------------------------------------
    # Save chunks
    # --------------------------------------------------

    with open(CHUNKS_FILE, "wb") as file:

        pickle.dump(
            chunks,
            file
        )


    print(
        f"Chunks saved: {CHUNKS_FILE}"
    )


    # --------------------------------------------------
    # Create FAISS vector database
    # --------------------------------------------------

    print("\nCreating FAISS database...")

    embeddings = HuggingFaceEmbeddings(
        model_name="all-MiniLM-L6-v2"
    )

    vectorstore = FAISS.from_documents(
        chunks,
        embeddings
    )


    # --------------------------------------------------
    # Save FAISS
    # --------------------------------------------------

    vectorstore.save_local(
        str(FAISS_DB_FOLDER)
    )

    print(
        "FAISS database created successfully!"
    )


    # --------------------------------------------------
    # Final status
    # --------------------------------------------------

    # print("\n========================================")
    # print("INDEXING COMPLETED")
    # print("========================================")
    # print(f"PDF pages : {len(documents)}")
    # print(f"Chunks    : {len(chunks)}")
    # print(f"FAISS     : {FAISS_DB_FOLDER}/")
    # print(f"BM25      : {BM25_FILE}")
    # print(f"Chunks    : {CHUNKS_FILE}")
    # print("========================================")

    print("========================================")
    print("INDEXING COMPLETED")
    print("========================================")
    print(f"Total documents : {len(documents)}")
    print(f"Total chunks    : {len(chunks)}")
    print(f"FAISS           : faiss_db/")
    print(f"BM25            : bm25.pkl")
    print(f"Chunks          : chunks.pkl")
    print("========================================")



if __name__ == "__main__":
    create_vector_db()
