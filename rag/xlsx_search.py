from pathlib import Path

from openpyxl import load_workbook


DATA_FOLDER = Path("data")


def clean_header(value):
    """
    Convert a cell value into a clean header name.
    """

    if value is None:
        return ""

    return str(value).strip()


def is_valid_header_row(row):
    """
    Check whether a row looks like a real table header.
    """

    headers = [
        clean_header(value)
        for value in row
        if clean_header(value)
    ]

    # A real table should have at least 3 columns.
    if len(headers) < 3:
        return False

    # Reject obvious title rows.
    title_words = [
        "summary",
        "roadmap",
        "learning roadmap",
        "automation planning"
    ]

    combined_text = " ".join(headers).lower()

    if any(
        title_word in combined_text
        for title_word in title_words
    ):
        return False

    return True


def find_header_row(worksheet, max_scan_rows=15):
    """
    Find the actual header row.

    Some Excel sheets have title rows or blank rows
    before the real table header.
    """

    for row_number, row in enumerate(
        worksheet.iter_rows(
            min_row=1,
            max_row=max_scan_rows,
            values_only=True
        ),
        start=1
    ):

        if is_valid_header_row(row):
            return row_number

    return None


def load_xlsx_files():
    """
    Load all XLSX files from the data folder.

    Automatically detects the real header row
    for each worksheet.
    """

    rows = []

    xlsx_files = list(DATA_FOLDER.glob("*.xlsx"))

    print(f"XLSX files found: {len(xlsx_files)}")

    for xlsx_path in xlsx_files:

        print(f"\nReading XLSX: {xlsx_path.name}")

        workbook = load_workbook(
            xlsx_path,
            read_only=True,
            data_only=True
        )

        for worksheet in workbook.worksheets:

            print(f"\n  Sheet: {worksheet.title}")

            # ---------------------------------------------
            # Find actual header row
            # ---------------------------------------------

            header_row_number = find_header_row(
                worksheet
            )

            if header_row_number is None:

                print(
                    "  No valid header row found. Skipping."
                )

                continue

            print(
                f"  Header row: {header_row_number}"
            )

            # ---------------------------------------------
            # Read header row
            # ---------------------------------------------

            header_row = next(
                worksheet.iter_rows(
                    min_row=header_row_number,
                    max_row=header_row_number,
                    values_only=True
                )
            )

            headers = [
                clean_header(value)
                for value in header_row
            ]

            # Keep only non-empty headers
            valid_headers = [
                (index, header)
                for index, header in enumerate(headers)
                if header
            ]

            if len(valid_headers) < 3:

                print(
                    "  Not enough valid headers. Skipping."
                )

                continue

            print(
                "  Columns:",
                [
                    header
                    for _, header in valid_headers
                ]
            )

            # ---------------------------------------------
            # Read data rows
            # ---------------------------------------------

            for row_number, row_values in enumerate(
                worksheet.iter_rows(
                    min_row=header_row_number + 1,
                    values_only=True
                ),
                start=header_row_number + 1
            ):

                row = {}

                for index, header in valid_headers:

                    value = (
                        row_values[index]
                        if index < len(row_values)
                        else None
                    )

                    row[header] = (
                        str(value).strip()
                        if value is not None
                        else ""
                    )

                # -----------------------------------------
                # Skip completely empty rows
                # -----------------------------------------

                has_data = any(
                    value.strip()
                    for key, value in row.items()
                    if not key.startswith("_")
                )

                if not has_data:
                    continue

                # -----------------------------------------
                # Metadata
                # -----------------------------------------

                row["_source"] = xlsx_path.name

                row["_file_type"] = "xlsx"

                row["_sheet"] = worksheet.title

                row["_page"] = row_number

                rows.append(row)

        workbook.close()

    print(
        f"\nTotal XLSX rows loaded: {len(rows)}"
    )

    return rows