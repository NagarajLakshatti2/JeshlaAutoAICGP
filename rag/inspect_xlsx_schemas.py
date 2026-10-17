from pathlib import Path
from openpyxl import load_workbook


DATA_FOLDER = Path("data")


xlsx_files = list(DATA_FOLDER.glob("*.xlsx"))


print("\n" + "=" * 80)
print("XLSX SCHEMA INSPECTION")
print("=" * 80)


for xlsx_path in xlsx_files:

    print("\n" + "-" * 80)
    print(f"FILE: {xlsx_path.name}")
    print("-" * 80)

    workbook = load_workbook(
        xlsx_path,
        read_only=True,
        data_only=True
    )

    for worksheet in workbook.worksheets:

        print(f"\nSHEET: {worksheet.title}")

        rows = worksheet.iter_rows(
            values_only=True
        )

        try:
            headers = next(rows)
        except StopIteration:
            print("Empty sheet")
            continue

        print("\nHEADERS:")

        for index, header in enumerate(headers, start=1):
            print(f"  {index}. {header!r}")

        print("\nFIRST 3 DATA ROWS:")

        for row_number, row in enumerate(rows, start=2):

            print(f"\nRow {row_number}:")

            for index, value in enumerate(row, start=1):

                if value is not None:
                    print(f"  Column {index}: {value!r}")

            if row_number >= 4:
                break

    workbook.close()